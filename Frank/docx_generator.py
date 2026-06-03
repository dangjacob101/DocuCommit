import json
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY
}


def _apply_paragraph_attrs(p, node):
    align = node.get("attrs", {}).get("textAlign")
    if align in ALIGN_MAP:
        p.alignment = ALIGN_MAP[align]


def _add_runs_to_paragraph(p, node_content):
    if not node_content:
        return
    for child in node_content:
        if child.get("type") == "text":
            text = child.get("text", "")
            run = p.add_run(text)
            for mark in child.get("marks", []):
                mtype = mark.get("type")
                if mtype == "bold":
                    run.bold = True
                elif mtype == "italic":
                    run.italic = True
                elif mtype == "underline":
                    run.underline = True
                elif mtype == "strike":
                    run.font.strike = True
                elif mtype == "superscript":
                    run.font.superscript = True
                elif mtype == "subscript":
                    run.font.subscript = True


def _process_block(doc, node):
    node_type = node.get("type")

    if node_type == "paragraph":
        p = doc.add_paragraph()
        _apply_paragraph_attrs(p, node)
        _add_runs_to_paragraph(p, node.get("content", []))

    elif node_type == "heading":
        level = node.get("attrs", {}).get("level", 1)
        p = doc.add_paragraph()
        p.style = f"Heading {level}"
        _apply_paragraph_attrs(p, node)
        _add_runs_to_paragraph(p, node.get("content", []))

    elif node_type == "blockquote":
        for child in node.get("content", []):
            if child.get("type") == "paragraph":
                p = doc.add_paragraph(style="Quote")
                _apply_paragraph_attrs(p, child)
                _add_runs_to_paragraph(p, child.get("content", []))
            else:
                _process_block(doc, child)

    elif node_type == "bulletList":
        for item in node.get("content", []):
            if item.get("type") == "listItem":
                for child in item.get("content", []):
                    if child.get("type") == "paragraph":
                        p = doc.add_paragraph(style="List Bullet")
                        _apply_paragraph_attrs(p, child)
                        _add_runs_to_paragraph(p, child.get("content", []))
                    else:
                        _process_block(doc, child)

    elif node_type == "orderedList":
        for item in node.get("content", []):
            if item.get("type") == "listItem":
                for child in item.get("content", []):
                    if child.get("type") == "paragraph":
                        p = doc.add_paragraph(style="List Number")
                        _apply_paragraph_attrs(p, child)
                        _add_runs_to_paragraph(p, child.get("content", []))
                    else:
                        _process_block(doc, child)

    elif node_type == "horizontalRule":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("―" * 40)
        run.font.color.rgb = RGBColor(128, 128, 128)

    elif node_type == "codeBlock":
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        for child in node.get("content", []):
            if child.get("type") == "text":
                text = child.get("text", "")
                run = p.add_run(text)
                run.font.name = "Courier New"


def generate_docx_from_content(title: str, content: str) -> BytesIO:
    doc = Document()

    # Page setup - standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Add title as main heading
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    is_json = False
    if content:
        try:
            data = json.loads(content)
            if isinstance(data, dict) and data.get("type") == "doc":
                is_json = True
                for block in data.get("content", []):
                    _process_block(doc, block)
        except (ValueError, TypeError):
            pass

    if not is_json:
        # Fallback to plain text line-by-line
        fallback_text = content or ""
        for line in fallback_text.splitlines():
            clean = line.strip()
            if clean:
                doc.add_paragraph(clean)

    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
