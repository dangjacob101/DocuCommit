import os
import sys
import unittest
from io import BytesIO

os.environ["DATABASE_URL"] = "sqlite:///:memory:"
sys.path.insert(0, os.path.dirname(__file__))

from app import create_app
from models import db, Branch, Commit
from docx import Document


class ExportRouteTest(unittest.TestCase):
    def setUp(self):
        self.app = create_app()
        self.app.config.update(TESTING=True)
        self.client = self.app.test_client()
        with self.app.app_context():
            db.drop_all()
            db.create_all()
            from app import _seed_default_user
            _seed_default_user()
        self._login()

    def _login(self):
        self.client.post(
            "/api/auth/login",
            json={"email": "frank@docucommit.com", "password": "CS35LTeamprofile!"},
        )

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()

    def test_export_docx_formatting(self):
        # Create a document with complex ProseMirror formatting JSON
        doc_content = {
            "type": "doc",
            "content": [
                {
                    "type": "heading",
                    "attrs": {"level": 1},
                    "content": [{"type": "text", "text": "Project Outline"}],
                },
                {
                    "type": "paragraph",
                    "attrs": {"textAlign": "center"},
                    "content": [
                        {"type": "text", "text": "This is a centered paragraph with "},
                        {
                            "type": "text",
                            "marks": [{"type": "bold"}, {"type": "italic"}],
                            "text": "bold italic runs",
                        },
                    ],
                },
                {
                    "type": "bulletList",
                    "content": [
                        {
                            "type": "listItem",
                            "content": [
                                {
                                    "type": "paragraph",
                                    "content": [{"type": "text", "text": "First bullet item"}],
                                }
                            ],
                        }
                    ],
                },
            ],
        }
        import json
        json_str = json.dumps(doc_content)

        r_create = self.client.post(
            "/api/documents", json={"title": "Design Specs", "content": json_str}
        )
        self.assertEqual(r_create.status_code, 201)
        doc = r_create.get_json()

        # Call export endpoint
        response = self.client.get(f"/api/documents/{doc['id']}/export")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        # Parse response using python-docx to verify content
        docx_file = Document(BytesIO(response.data))

        # Check document title heading (Heading 0)
        self.assertEqual(docx_file.paragraphs[0].text, "Design Specs")

        # Check Heading 1 ("Project Outline")
        self.assertEqual(docx_file.paragraphs[1].text, "Project Outline")
        self.assertEqual(docx_file.paragraphs[1].style.name, "Heading 1")

        # Check paragraph text and runs
        center_para = docx_file.paragraphs[2]
        self.assertEqual(
            center_para.text, "This is a centered paragraph with bold italic runs"
        )
        # Check alignment (centered is 1 in Word)
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        self.assertEqual(center_para.alignment, WD_ALIGN_PARAGRAPH.CENTER)

        # Check bold/italic marks on the second run
        # Note: docx might split runs slightly differently, but let's check bold/italic run exists
        bold_italic_runs = [
            r for r in center_para.runs if r.bold and r.italic
        ]
        self.assertTrue(len(bold_italic_runs) > 0)
        self.assertEqual(bold_italic_runs[0].text, "bold italic runs")

        # Check list item
        self.assertEqual(docx_file.paragraphs[3].text, "First bullet item")
        self.assertEqual(docx_file.paragraphs[3].style.name, "List Bullet")


if __name__ == "__main__":
    unittest.main()
